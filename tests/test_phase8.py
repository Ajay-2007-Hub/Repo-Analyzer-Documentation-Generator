from unittest.mock import Mock

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import ValidationError

from app.generators.ppt import PowerPointGenerator
from app.models.architecture import ArchEdge, ArchitectureGraph, ArchNode
from app.models.presentation import SLIDE_TITLES, PresentationContent
from app.models.schemas import (
    AuthInfo,
    DatabaseInfo,
    Endpoint,
    ExternalService,
    FileInfo,
    ProjectAnalysis,
    RankedFile,
    RepoInfo,
    ScanResult,
    SecurityFinding,
)


def sample_inputs():
    scan_result = ScanResult(
        repository_name="sample-repo",
        commit_sha="abc123",
        total_files=3,
        total_lines=100,
        total_size_mb=0.5,
        files=[
            FileInfo(
                path="src/main.py",
                size_kb=2.5,
                lines=60,
                extension=".py",
                is_binary=False,
            )
        ],
        file_tree={},
    )
    project_analysis = ProjectAnalysis(
        repo_info=RepoInfo(
            languages={"Python": 100.0},
            frameworks=["FastAPI"],
            dependencies=["fastapi", "pydantic"],
            entry_points=["src/main.py"],
        ),
        top_files=[
            RankedFile(path="src/main.py", score=100, reason="Entry point")
        ],
    )
    endpoints = [
        Endpoint(
            method="GET",
            path="/health",
            file="src/main.py",
            description="Check service health",
        )
    ]
    architecture = ArchitectureGraph(
        nodes=[
            ArchNode(
                id="api",
                label="FastAPI",
                node_type="backend",
                description="HTTP service",
            ),
            ArchNode(
                id="db",
                label="SQLite",
                node_type="database",
                description="Persistent data",
            ),
        ],
        edges=[
            ArchEdge(
                from_node="api",
                to_node="db",
                label="query",
                edge_type="data",
            )
        ],
    )
    return scan_result, project_analysis, endpoints, architecture


def sample_content():
    return PresentationContent(
        slides=[
            {
                "title": title,
                "bullets": [
                    f"{title} evidence point one",
                    f"{title} evidence point two",
                    f"{title} evidence point three",
                ],
            }
            for title in SLIDE_TITLES
        ]
    )


def test_presentation_schema_enforces_exact_titles_and_bullet_limits():
    content = sample_content()
    assert tuple(slide.title for slide in content.slides) == SLIDE_TITLES
    assert all(3 <= len(slide.bullets) <= 5 for slide in content.slides)

    bad_content = content.model_dump()
    bad_content["slides"][0]["bullets"][0] = "word " * 13
    with pytest.raises(ValidationError, match="at most 12 words"):
        PresentationContent.model_validate(bad_content)

    bad_content = content.model_dump()
    bad_content["slides"][1]["title"] = "Wrong title"
    with pytest.raises(ValidationError):
        PresentationContent.model_validate(bad_content)


def test_generator_builds_and_reopens_exact_ten_slides(tmp_path):
    scan, analysis, endpoints, architecture = sample_inputs()
    generator = PowerPointGenerator(output_dir=tmp_path)
    output = generator.generate(
        scan,
        analysis,
        endpoints,
        db_info=DatabaseInfo(db_type="SQLite", orm="SQLAlchemy"),
        auth_info=AuthInfo(mechanism="JWT"),
        external_services=[ExternalService(domain="api.example.test")],
        security_findings=[
            SecurityFinding(
                category="Configuration",
                severity="MEDIUM",
                file="src/main.py",
                line=1,
                description="Debug mode is enabled",
            )
        ],
        architecture=architecture,
        slide_content=sample_content(),
    )

    reopened = Presentation(str(output))
    assert len(reopened.slides) == 10
    assert output.name == "sample-repo_presentation.pptx"
    assert reopened.core_properties.author == "RepoInsight AI"

    cover_text = "\n".join(
        shape.text for shape in reopened.slides[0].shapes if shape.has_text_frame
    )
    assert "sample-repo" in cover_text
    for index, title in enumerate(SLIDE_TITLES):
        slide = reopened.slides[index]
        for shape in slide.shapes:
            assert shape.left >= 0
            assert shape.top >= 0
            assert shape.left + shape.width <= reopened.slide_width
            assert shape.top + shape.height <= reopened.slide_height
        all_text = "\n".join(
            shape.text for shape in slide.shapes if shape.has_text_frame
        )
        if index:
            assert title in all_text
        bullet_paragraphs = [
            paragraph
            for shape in slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            if paragraph._p.pPr is not None
            and paragraph._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
            is not None
        ]
        assert len(bullet_paragraphs) in (3, 4, 5)
        assert all(
            paragraph.font.size is None
            or paragraph.font.size.pt <= 22
            for paragraph in bullet_paragraphs
        )

    architecture_slide = reopened.slides[5]
    assert any(
        shape.name.startswith("Architecture node:") for shape in architecture_slide.shapes
    )
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.LINE
        for shape in architecture_slide.shapes
    )


def test_generator_uses_ollama_with_compact_evidence_prompt(tmp_path):
    scan, analysis, endpoints, architecture = sample_inputs()
    llm = Mock()
    llm.generate_structured.return_value = sample_content()
    generator = PowerPointGenerator(tmp_path, ollama_client=llm)

    output = generator.generate(
        scan,
        analysis,
        endpoints,
        db_info=DatabaseInfo(db_type="SQLite"),
        architecture=architecture,
    )

    assert output.is_file()
    llm.generate_structured.assert_called_once()
    prompt, schema = llm.generate_structured.call_args.args[:2]
    assert schema is PresentationContent
    assert "Use only the evidence below" in prompt
    assert "GET /health" in prompt
    assert "src/main.py" in prompt
    assert "12 words" in prompt
    assert "repository code" not in prompt.lower()
    assert llm.generate_structured.call_args.kwargs["retries"] == 2


def test_generator_embeds_supplied_diagram_and_honors_blank_template(tmp_path):
    from PIL import Image

    scan, analysis, endpoints, architecture = sample_inputs()
    diagram_path = tmp_path / "architecture.png"
    Image.new("RGB", (640, 240), color=(230, 240, 245)).save(diagram_path)

    blank_template = Presentation()
    template_path = tmp_path / "blank_template.pptx"
    blank_template.save(template_path)

    output = PowerPointGenerator(
        output_dir=tmp_path, template_path=template_path
    ).generate(
        scan,
        analysis,
        endpoints,
        architecture=architecture,
        diagram_path=diagram_path,
        slide_content=sample_content(),
    )
    reopened = Presentation(str(output))
    architecture_slide = reopened.slides[5]
    assert any(
        shape.shape_type == 13 and shape.name == "Architecture diagram"
        for shape in architecture_slide.shapes
    )
