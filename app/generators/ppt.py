import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.llm.ollama_client import OllamaClient
from app.models.architecture import ArchitectureGraph
from app.models.presentation import SLIDE_TITLES, PresentationContent
from app.models.report import ReportNarrative
from app.models.schemas import (
    AuthInfo,
    DatabaseInfo,
    Endpoint,
    ExternalService,
    ProjectAnalysis,
    ScanResult,
    SecurityFinding,
)

NAVY = RGBColor(24, 59, 78)
TEAL = RGBColor(19, 138, 138)
PALE = RGBColor(242, 247, 249)
INK = RGBColor(43, 58, 69)
MUTED = RGBColor(100, 118, 131)
WHITE = RGBColor(255, 255, 255)
NODE_COLORS = {
    "frontend": RGBColor(225, 241, 246),
    "backend": RGBColor(235, 226, 245),
    "database": RGBColor(247, 237, 222),
    "auth": RGBColor(225, 243, 229),
    "external": RGBColor(248, 229, 229),
    "module": RGBColor(235, 240, 243),
}


class PowerPointGenerator:
    """Create a ten-slide project presentation from validated analysis facts."""

    def __init__(
        self,
        output_dir: str | Path = "outputs",
        template_path: str | Path | None = None,
        ollama_client: OllamaClient | None = None,
    ):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.template_path = Path(template_path) if template_path else None
        self.ollama_client = ollama_client or OllamaClient()

    def generate(
        self,
        scan_result: ScanResult,
        project_analysis: ProjectAnalysis,
        endpoints: list[Endpoint],
        db_info: DatabaseInfo | None = None,
        auth_info: AuthInfo | None = None,
        external_services: list[ExternalService] | None = None,
        security_findings: list[SecurityFinding] | None = None,
        architecture: ArchitectureGraph | None = None,
        diagram_path: str | Path | None = None,
        narrative: ReportNarrative | None = None,
        slide_content: PresentationContent | None = None,
        output_path: str | Path | None = None,
    ) -> Path:
        """Generate and save the deck, using Ollama unless validated slides are supplied."""
        if slide_content is None:
            prompt = self._build_prompt(
                scan_result=scan_result,
                project_analysis=project_analysis,
                endpoints=endpoints,
                db_info=db_info,
                auth_info=auth_info,
                external_services=external_services or [],
                security_findings=security_findings or [],
                architecture=architecture,
                narrative=narrative,
            )
            slide_content = self.ollama_client.generate_structured(
                prompt, PresentationContent, retries=2
            )

        if self.template_path:
            if not self.template_path.is_file():
                raise FileNotFoundError(f"PowerPoint template not found: {self.template_path}")
            presentation = Presentation(str(self.template_path))
        else:
            presentation = Presentation()

        if len(presentation.slides):
            raise ValueError(
                "PowerPoint template must not contain slides; use a blank template."
            )

        self._configure_presentation(presentation)
        for index, slide_spec in enumerate(slide_content.slides):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if index == 0:
                self._add_cover(slide, slide_spec, scan_result.repository_name)
            else:
                self._add_content_slide(
                    slide,
                    slide_spec.title,
                    slide_spec.bullets,
                    slide_number=index + 1,
                    architecture=architecture if slide_spec.title == "Architecture" else None,
                    diagram_path=diagram_path
                    if slide_spec.title == "Architecture"
                    else None,
                )

        safe_name = re.sub(
            r"[^A-Za-z0-9_-]+", "_", scan_result.repository_name
        ).strip("_")
        destination = (
            Path(output_path)
            if output_path is not None
            else self.output_dir / f"{safe_name or 'repository'}_presentation.pptx"
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(destination))
        return destination

    def _build_prompt(
        self,
        *,
        scan_result: ScanResult,
        project_analysis: ProjectAnalysis,
        endpoints: list[Endpoint],
        db_info: DatabaseInfo | None,
        auth_info: AuthInfo | None,
        external_services: list[ExternalService],
        security_findings: list[SecurityFinding],
        architecture: ArchitectureGraph | None,
        narrative: ReportNarrative | None,
    ) -> str:
        info = project_analysis.repo_info
        endpoint_samples = sorted(
            {f"{item.method} {item.path}" for item in endpoints}
        )[:12]
        evidence = [
            f"Repository: {scan_result.repository_name}",
            f"Commit: {scan_result.commit_sha or 'Not detected'}",
            f"Files: {scan_result.total_files}; lines: {scan_result.total_lines}",
            f"Languages: {self._limit_list([f'{name} {percent:.1f}%' for name, percent in info.languages.items()], 12)}",
            f"Frameworks: {self._limit_list(info.frameworks, 12)}",
            f"Dependencies: {self._limit_list(sorted(set(info.dependencies)), 30)}",
            f"Entry points: {self._limit_list(info.entry_points, 10)}",
            f"Endpoints detected: {len(endpoints)}; examples: {self._limit_list(endpoint_samples, 12)}",
            f"Database: {db_info.db_type if db_info else 'Not detected'}",
            f"ORM: {db_info.orm if db_info and db_info.orm else 'Not detected'}",
            f"Authentication: {auth_info.mechanism if auth_info else 'Not detected'}",
            "External domains: "
            + self._limit_list(
                sorted({service.domain for service in external_services}), 12
            ),
            "Security finding categories: "
            + self._limit_list(
                sorted({finding.category for finding in security_findings}), 12
            ),
            "Important files: "
            + self._limit_list(
                [
                    f"{item.path} ({item.reason})"
                    for item in project_analysis.top_files[:10]
                ],
                10,
            ),
            "Architecture components: "
            + self._limit_list(
                [
                    f"{node.label} [{node.node_type}]"
                    for node in architecture.nodes[:12]
                ]
                if architecture
                else [],
                12,
            ),
        ]
        if architecture and architecture.description:
            evidence.append(f"Architecture summary: {architecture.description[:600]}")
        if narrative:
            if narrative.problem_statement:
                evidence.append(
                    f"Validated problem statement: {narrative.problem_statement[:400]}"
                )
            evidence.extend(
                [
                    f"Validated objective: {item[:160]}"
                    for item in narrative.objectives[:5]
                ]
            )
            evidence.extend(
                [
                    f"Validated workflow: {item[:160]}"
                    for item in narrative.working_flow[:5]
                ]
            )
            evidence.extend(
                [
                    f"Validated future enhancement: {item[:160]}"
                    for item in narrative.future_enhancements[:5]
                ]
            )

        return (
            "Create concise slide bullets for a repository analysis presentation.\n"
            "Use only the evidence below. If evidence is missing, say 'Not detected'. "
            "Do not invent files, APIs, technologies, business goals, or architecture.\n"
            "Return JSON matching the supplied schema with exactly 10 slides in this "
            f"order: {', '.join(SLIDE_TITLES)}.\n"
            "Every slide must contain 3-5 bullets. Every bullet must contain at most "
            "12 words. Keep bullets factual, clear, and presentation-ready. Do not "
            "include secrets or credentials.\n\n"
            "Evidence:\n- "
            + "\n- ".join(evidence)
        )

    @staticmethod
    def _limit_list(values: list[str], limit: int) -> str:
        if not values:
            return "Not detected"
        text = ", ".join(values[:limit])
        if len(values) > limit:
            text += f", and {len(values) - limit} more"
        return text

    @staticmethod
    def _configure_presentation(presentation: Presentation) -> None:
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = "RepoInsight AI Analysis"
        presentation.core_properties.author = "RepoInsight AI"
        presentation.core_properties.subject = "Repository analysis presentation"

    def _add_cover(self, slide, slide_spec, repo_name: str) -> None:
        self._set_background(slide, NAVY)
        self._add_rectangle(slide, 0, 0, 0.24, 7.5, TEAL)
        self._add_text(
            slide,
            "REPOINSIGHT AI",
            0.85,
            0.75,
            5.2,
            0.4,
            font_size=15,
            color=TEAL,
            bold=True,
        )
        self._add_text(
            slide,
            repo_name,
            0.85,
            1.45,
            11.4,
            1.3,
            font_size=34,
            color=WHITE,
            bold=True,
        )
        self._add_text(
            slide,
            "Repository Analysis",
            0.88,
            2.8,
            10.8,
            0.55,
            font_size=22,
            color=RGBColor(189, 215, 224),
        )
        self._add_bullet_box(
            slide,
            slide_spec.bullets,
            0.88,
            4.1,
            11.25,
            2.25,
            font_size=18,
            color=WHITE,
        )
        self._add_text(
            slide,
            "1 / 10",
            11.75,
            7.04,
            0.8,
            0.2,
            font_size=9,
            color=RGBColor(189, 215, 224),
            align=PP_ALIGN.RIGHT,
        )

    def _add_content_slide(
        self,
        slide,
        title: str,
        bullets: list[str],
        slide_number: int,
        architecture: ArchitectureGraph | None,
        diagram_path: str | Path | None,
    ) -> None:
        self._set_background(slide, PALE)
        self._add_text(
            slide,
            "REPOINSIGHT AI  /  REPOSITORY ANALYSIS",
            0.72,
            0.34,
            8,
            0.23,
            font_size=9,
            color=TEAL,
            bold=True,
        )
        self._add_text(
            slide,
            title,
            0.72,
            0.72,
            11.8,
            0.62,
            font_size=27,
            color=NAVY,
            bold=True,
        )
        self._add_rectangle(slide, 0.72, 1.42, 0.72, 0.055, TEAL)
        if title == "Architecture":
            self._add_bullet_box(
                slide, bullets, 0.82, 1.7, 11.6, 1.65, font_size=17
            )
            self._add_architecture_visual(
                slide, architecture, diagram_path, 0.82, 3.58, 11.6, 2.95
            )
        else:
            self._add_bullet_box(
                slide, bullets, 0.88, 1.8, 11.55, 4.85, font_size=21
            )
        self._add_rectangle(slide, 0.72, 7.03, 11.9, 0.012, RGBColor(214, 225, 231))
        self._add_text(
            slide,
            f"{slide_number} / 10",
            11.73,
            7.08,
            0.78,
            0.2,
            font_size=9,
            color=MUTED,
            align=PP_ALIGN.RIGHT,
        )

    def _add_architecture_visual(
        self,
        slide,
        architecture: ArchitectureGraph | None,
        diagram_path: str | Path | None,
        x: float,
        y: float,
        width: float,
        height: float,
    ) -> None:
        if diagram_path is not None:
            path = Path(diagram_path)
            if not path.is_file():
                raise FileNotFoundError(f"Architecture diagram not found: {path}")
            from PIL import Image as PILImage

            with PILImage.open(path) as image:
                image_width, image_height = image.size
            scale = min(width / image_width, height / image_height)
            picture = slide.shapes.add_picture(
                str(path),
                Inches(x + (width - image_width * scale) / 2),
                Inches(y + (height - image_height * scale) / 2),
                width=Inches(image_width * scale),
                height=Inches(image_height * scale),
            )
            picture.name = "Architecture diagram"
            return

        if architecture and architecture.nodes:
            self._draw_graph(slide, architecture, x, y, width, height)
            return

        self._add_text(
            slide,
            "Architecture diagram not detected",
            x + 0.2,
            y + 1.0,
            width - 0.4,
            0.5,
            font_size=18,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )

    def _draw_graph(
        self,
        slide,
        architecture: ArchitectureGraph,
        x: float,
        y: float,
        width: float,
        height: float,
    ) -> None:
        nodes = architecture.nodes[:12]
        columns = min(4, max(1, len(nodes)))
        rows = math.ceil(len(nodes) / columns)
        cell_width = width / columns
        cell_height = height / rows
        node_width = min(2.45, cell_width * 0.86)
        node_height = min(0.68, cell_height * 0.66)
        positions = {}

        for index, node in enumerate(nodes):
            row, column = divmod(index, columns)
            center_x = x + column * cell_width + cell_width / 2
            center_y = y + row * cell_height + cell_height / 2
            positions[node.id] = (center_x, center_y)

        for edge in architecture.edges:
            if edge.from_node not in positions or edge.to_node not in positions:
                continue
            start_x, start_y = positions[edge.from_node]
            end_x, end_y = positions[edge.to_node]
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(start_x),
                Inches(start_y),
                Inches(end_x),
                Inches(end_y),
            )
            connector.line.color.rgb = MUTED
            connector.line.width = Pt(1.25)
            if edge.label:
                self._add_text(
                    slide,
                    edge.label,
                    (start_x + end_x) / 2 - 0.5,
                    (start_y + end_y) / 2 - 0.1,
                    1,
                    0.2,
                    font_size=8,
                    color=MUTED,
                    align=PP_ALIGN.CENTER,
                )

        for index, node in enumerate(nodes):
            row, column = divmod(index, columns)
            left = x + column * cell_width + (cell_width - node_width) / 2
            top = y + row * cell_height + (cell_height - node_height) / 2
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(left),
                Inches(top),
                Inches(node_width),
                Inches(node_height),
            )
            shape.name = f"Architecture node: {node.label}"
            shape.fill.solid()
            shape.fill.fore_color.rgb = NODE_COLORS.get(node.node_type, PALE)
            shape.line.color.rgb = RGBColor(124, 151, 166)
            shape.text_frame.clear()
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            shape.text_frame.word_wrap = True
            shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            shape.text_frame.margin_left = Inches(0.08)
            shape.text_frame.margin_right = Inches(0.08)
            shape.text_frame.margin_top = Inches(0.03)
            shape.text_frame.margin_bottom = Inches(0.03)
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.text = node.label
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(13)
            paragraph.font.bold = True
            paragraph.font.color.rgb = NAVY

    @staticmethod
    def _set_background(slide, color: RGBColor) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color

    @staticmethod
    def _add_rectangle(
        slide, x: float, y: float, width: float, height: float, color: RGBColor
    ) -> None:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    def _add_text(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        width: float,
        height: float,
        *,
        font_size: int,
        color: RGBColor,
        bold: bool = False,
        align=PP_ALIGN.LEFT,
    ):
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        shape.text_frame.clear()
        shape.text_frame.word_wrap = True
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.text_frame.margin_left = 0
        shape.text_frame.margin_right = 0
        shape.text_frame.margin_top = 0
        shape.text_frame.margin_bottom = 0
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = align
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return shape

    def _add_bullet_box(
        self,
        slide,
        bullets: list[str],
        x: float,
        y: float,
        width: float,
        height: float,
        *,
        font_size: int,
        color: RGBColor = INK,
    ) -> None:
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        shape.text_frame.clear()
        shape.text_frame.word_wrap = True
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.text_frame.margin_left = Inches(0.08)
        shape.text_frame.margin_right = Inches(0.08)
        shape.text_frame.margin_top = Inches(0.08)
        shape.text_frame.margin_bottom = Inches(0.08)
        for index, bullet in enumerate(bullets):
            paragraph = (
                shape.text_frame.paragraphs[0]
                if index == 0
                else shape.text_frame.add_paragraph()
            )
            paragraph.text = bullet
            paragraph.level = 0
            paragraph_properties = paragraph._p.get_or_add_pPr()
            bullet_character = OxmlElement("a:buChar")
            bullet_character.set("char", "•")
            paragraph_properties.insert(0, bullet_character)
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = color
            paragraph.space_after = Pt(12 if len(bullets) < 5 else 7)
            paragraph.left_indent = Inches(0.28)
            paragraph.first_line_indent = Inches(-0.22)
        shape.text_frame.word_wrap = True
